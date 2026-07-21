import argparse
import contextlib
import datetime
import hashlib
import json
import os
import re
import sqlite3
import time

import fitz

from model_client import ModelClient

DATE_DIR_PATTERN = re.compile(r"\d{4}-\d{2}-\d{2}")
SUPPORTED_EXTS = {".pdf", ".docx", ".doc", ".md", ".txt"}
SKIP_SUFFIXES = ("_summary.md",)
PARSER_VERSION = "mupdf-text-v1"
CHUNKER_VERSION = "sections-v1-chunks-300-600-overlap-15"

DEFAULT_SOURCE_DIR = os.path.join(
    os.getcwd(), datetime.datetime.now().strftime("%Y-%m-%d"), "HuggingFace"
)
DEFAULT_KB_DIR = os.path.join(os.getcwd(), "kb_store")
DEFAULT_API_KEY_PATH = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "API_KEY.json")
)


class CorruptIndexError(RuntimeError):
    pass


class EmbeddingModelChangedError(RuntimeError):
    """Raised when the configured embedding model differs from the one used to build the index.

    The stored vectors are tied to the old embedding model, but the chunk text in
    doc_store.jsonl can be re-embedded with the new model, so this is recoverable.
    """

    def __init__(self, old_model, new_model):
        self.old_model = old_model
        self.new_model = new_model
        super().__init__(
            f"Embedding model changed ({old_model or 'unknown'} -> {new_model})."
        )


def quarantine_sqlite_file(path):
    stamp = int(time.time())
    for candidate in (path, f"{path}-wal", f"{path}-shm"):
        if os.path.exists(candidate):
            with contextlib.suppress(Exception):
                os.replace(candidate, f"{candidate}.corrupt.{stamp}")


class EmbeddingCache:
    """Durable normalized embeddings keyed by chunk content and model.

    FAISS is a derived index. Keeping the vectors separately means an index
    repair does not need to call the embedding service again.
    """

    def __init__(self, kb_dir, embed_model, np):
        model_key = hashlib.sha1(embed_model.encode("utf-8")).hexdigest()[:16]
        self.path = os.path.join(kb_dir, f"embedding_cache_{model_key}.sqlite3")
        self.np = np
        try:
            self.conn = sqlite3.connect(self.path, timeout=60)
            self.conn.execute("PRAGMA journal_mode=WAL")
            self.conn.execute("PRAGMA synchronous=FULL")
            self.conn.execute(
                "CREATE TABLE IF NOT EXISTS vectors ("
                "text_hash TEXT PRIMARY KEY, dimension INTEGER NOT NULL, "
                "checksum TEXT NOT NULL, vector BLOB NOT NULL)"
            )
            self.conn.execute(
                "CREATE TABLE IF NOT EXISTS metadata ("
                "key TEXT PRIMARY KEY, value TEXT NOT NULL)"
            )
            self.conn.commit()
        except sqlite3.DatabaseError:
            with contextlib.suppress(Exception):
                self.conn.close()
            quarantine_sqlite_file(self.path)
            self.conn = sqlite3.connect(self.path, timeout=60)
            self.conn.execute("PRAGMA journal_mode=WAL")
            self.conn.execute("PRAGMA synchronous=FULL")
            self.conn.execute(
                "CREATE TABLE IF NOT EXISTS vectors ("
                "text_hash TEXT PRIMARY KEY, dimension INTEGER NOT NULL, "
                "checksum TEXT NOT NULL, vector BLOB NOT NULL)"
            )
            self.conn.execute(
                "CREATE TABLE IF NOT EXISTS metadata ("
                "key TEXT PRIMARY KEY, value TEXT NOT NULL)"
            )
            self.conn.commit()

    def get_many(self, text_hashes):
        keys = list(dict.fromkeys(text_hashes))
        if not keys:
            return {}
        found = {}
        for start in range(0, len(keys), 800):
            batch = keys[start:start + 800]
            placeholders = ",".join("?" for _ in batch)
            rows = self.conn.execute(
                f"SELECT text_hash, dimension, checksum, vector FROM vectors "
                f"WHERE text_hash IN ({placeholders})",
                batch,
            )
            for text_hash, dimension, checksum, blob in rows:
                if hashlib.sha256(blob).hexdigest() != checksum:
                    continue
                vector = self.np.frombuffer(blob, dtype="float32").copy()
                if vector.size != int(dimension):
                    continue
                found[text_hash] = vector
        return found

    def put_many(self, vectors):
        if not vectors:
            return
        rows = []
        for text_hash, vector in vectors.items():
            vector = self.np.asarray(vector, dtype="float32")
            blob = vector.tobytes()
            rows.append(
                (text_hash, int(vector.size), hashlib.sha256(blob).hexdigest(), blob)
            )
        with self.conn:
            self.conn.executemany(
                "INSERT OR REPLACE INTO vectors(text_hash, dimension, checksum, vector) "
                "VALUES (?, ?, ?, ?)",
                rows,
            )

    def count(self):
        row = self.conn.execute("SELECT COUNT(*) FROM vectors").fetchone()
        return int(row[0]) if row else 0

    def seeded_index_size(self):
        row = self.conn.execute(
            "SELECT value FROM metadata WHERE key = 'seeded_index_size'"
        ).fetchone()
        if not row:
            return None
        try:
            return int(row[0])
        except (TypeError, ValueError):
            return None

    def mark_index_seeded(self, size):
        with self.conn:
            self.conn.execute(
                "INSERT OR REPLACE INTO metadata(key, value) VALUES (?, ?)",
                ("seeded_index_size", str(int(size))),
            )

    def close(self):
        with contextlib.suppress(Exception):
            self.conn.commit()
            self.conn.close()


class ChunkStore:
    """Canonical chunk records used to repair the JSONL compatibility store."""

    def __init__(self, kb_dir):
        self.path = os.path.join(kb_dir, "chunk_store.sqlite3")
        try:
            self.conn = sqlite3.connect(self.path, timeout=60)
            self.conn.execute("PRAGMA journal_mode=WAL")
            self.conn.execute("PRAGMA synchronous=FULL")
            self.conn.execute(
                "CREATE TABLE IF NOT EXISTS records ("
                "chunk_id TEXT PRIMARY KEY, vector_id INTEGER NOT NULL UNIQUE, "
                "record_json TEXT NOT NULL)"
            )
            self.conn.commit()
        except sqlite3.DatabaseError:
            with contextlib.suppress(Exception):
                self.conn.close()
            quarantine_sqlite_file(self.path)
            self.conn = sqlite3.connect(self.path, timeout=60)
            self.conn.execute("PRAGMA journal_mode=WAL")
            self.conn.execute("PRAGMA synchronous=FULL")
            self.conn.execute(
                "CREATE TABLE IF NOT EXISTS records ("
                "chunk_id TEXT PRIMARY KEY, vector_id INTEGER NOT NULL UNIQUE, "
                "record_json TEXT NOT NULL)"
            )
            self.conn.commit()

    def count(self):
        row = self.conn.execute("SELECT COUNT(*) FROM records").fetchone()
        return int(row[0]) if row else 0

    def ids_and_max(self):
        chunk_ids = set()
        vector_ids = set()
        max_vector_id = 0
        for chunk_id, vector_id in self.conn.execute("SELECT chunk_id, vector_id FROM records"):
            vector_id = int(vector_id)
            chunk_ids.add(chunk_id)
            vector_ids.add(vector_id)
            max_vector_id = max(max_vector_id, vector_id)
        return chunk_ids, vector_ids, max_vector_id

    def put_records(self, records):
        if not records:
            return
        rows = [
            (
                record["chunk_id"],
                int(record["vector_id"]),
                json.dumps(record, ensure_ascii=False),
            )
            for record in records
        ]
        with self.conn:
            self.conn.executemany(
                "INSERT OR REPLACE INTO records(chunk_id, vector_id, record_json) "
                "VALUES (?, ?, ?)",
                rows,
            )

    def seed_from_jsonl(self, doc_store_path):
        if not os.path.exists(doc_store_path) or self.count():
            return 0
        batch = []
        imported = 0
        for record in iter_doc_store_records(doc_store_path):
            batch.append(record)
            if len(batch) >= 512:
                self.put_records(batch)
                imported += len(batch)
                batch = []
        if batch:
            self.put_records(batch)
            imported += len(batch)
        return imported

    def iter_records(self):
        for (record_json,) in self.conn.execute(
            "SELECT record_json FROM records ORDER BY vector_id"
        ):
            try:
                record = json.loads(record_json)
            except Exception:
                continue
            if record.get("text") and isinstance(record.get("vector_id"), int):
                yield record

    def rewrite_jsonl(self, doc_store_path):
        tmp_path = f"{doc_store_path}.{os.getpid()}.tmp"
        try:
            with open(tmp_path, "w", encoding="utf-8") as f:
                for record in self.iter_records():
                    f.write(json.dumps(record, ensure_ascii=False) + "\n")
                f.flush()
                os.fsync(f.fileno())
            os.replace(tmp_path, doc_store_path)
        finally:
            if os.path.exists(tmp_path):
                with contextlib.suppress(Exception):
                    os.remove(tmp_path)

    def close(self):
        with contextlib.suppress(Exception):
            self.conn.commit()
            self.conn.close()


def chunk_text_hash(text):
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def emit(reporter, event, **payload):
    if reporter:
        reporter({"event": event, **payload})


def report_log(reporter, message, **payload):
    print(message)
    emit(reporter, "log", message=message, **payload)


def remove_corrupt_file(file_path, reason=None):
    base_name = os.path.basename(file_path)
    if not os.path.exists(file_path):
        print(f"[!] Corrupt file missing, skipped: {base_name}")
        return
    try:
        os.remove(file_path)
        print(f"[!] Corrupt file deleted: {base_name}")
    except Exception as exc:
        if reason:
            print(f"[!] Failed to delete corrupt file: {base_name} ({reason}) ({exc})")
        else:
            print(f"[!] Failed to delete corrupt file: {base_name} ({exc})")


def build_kb(source_dir=DEFAULT_SOURCE_DIR, kb_dir=DEFAULT_KB_DIR, api_key_path=DEFAULT_API_KEY_PATH, rebuild=False, refresh=False, reporter=None):
    try:
        import numpy as np
    except Exception:
        report_log(reporter, "[!] Missing dependency: numpy. Install it before running createbase.", stage="kb-build", level="error")
        return

    try:
        import faiss
    except Exception:
        report_log(reporter, "[!] Missing dependency: faiss-cpu. Install it before running createbase.", stage="kb-build", level="error")
        return

    if isinstance(source_dir, (list, tuple, set)):
        source_dir = [path for path in source_dir if os.path.exists(path)]
        source_exists = bool(source_dir)
    else:
        source_exists = bool(source_dir) and os.path.exists(source_dir)

    if not source_exists:
        if refresh:
            report_log(
                reporter,
                f"[*] Source directory not found ({source_dir}); refresh will re-embed existing chunks only.",
                stage="kb-build",
                level="warning",
            )
            source_dir = None
        else:
            report_log(reporter, f"[!] Source directory not found: {source_dir}", stage="kb-build", level="error")
            return
    emit(reporter, "stage", stage="kb-build", source_dir=source_dir, kb_dir=kb_dir)

    os.makedirs(kb_dir, exist_ok=True)
    build_lock = acquire_build_lock(kb_dir)
    index_path = os.path.join(kb_dir, "vectors.faiss")
    meta_path = os.path.join(kb_dir, "index_meta.json")
    doc_store_path = os.path.join(kb_dir, "doc_store.jsonl")
    manifest_path = os.path.join(kb_dir, "source_manifest.json")
    index_dirty = False
    meta_dirty = False

    if rebuild:
        if os.path.exists(index_path):
            os.remove(index_path)
        if os.path.exists(meta_path):
            os.remove(meta_path)
        if os.path.exists(doc_store_path):
            os.remove(doc_store_path)
        if os.path.exists(manifest_path):
            os.remove(manifest_path)
        remove_sqlite_store(os.path.join(kb_dir, "chunk_store.sqlite3"))
    elif refresh:
        # Refresh keeps the parsed chunk text (doc_store.jsonl) but drops the stale
        # FAISS vectors + meta so they get rebuilt with the current embedding model.
        if os.path.exists(index_path):
            os.remove(index_path)
        if os.path.exists(meta_path):
            os.remove(meta_path)
        report_log(
            reporter,
            "[*] Refresh requested: re-embedding existing chunks with the current embedding model.",
            stage="kb-build",
        )

    model_client = ModelClient.from_config_path(api_key_path)
    embed_model = model_client.embedding_conf.get("model")
    if not embed_model:
        print("[!] Embedding model is not configured in API_KEY.json.")
        build_lock.close()
        return

    chunk_store = ChunkStore(kb_dir)
    imported_chunks = chunk_store.seed_from_jsonl(doc_store_path)
    if imported_chunks:
        report_log(
            reporter,
            f"[*] Imported {imported_chunks} existing chunks into the durable chunk store.",
            stage="kb-build",
        )

    file_chunk_ids, file_max_vector_id, invalid_lines = inspect_doc_store(doc_store_path)
    existing_chunk_ids, stored_vector_ids, max_vector_id = chunk_store.ids_and_max()
    max_vector_id = max(max_vector_id, file_max_vector_id)
    if existing_chunk_ids and (
        invalid_lines or file_chunk_ids != existing_chunk_ids
    ):
        report_log(
            reporter,
            "[!] doc_store.jsonl is incomplete or damaged; restoring it from chunk_store.sqlite3.",
            stage="kb-build",
            level="warning",
        )
        chunk_store.rewrite_jsonl(doc_store_path)

    embedding_cache = EmbeddingCache(kb_dir, embed_model, np)
    source_manifest = load_source_manifest(manifest_path)

    try:
        index, meta = load_or_create_index(index_path, meta_path, embed_model)
    except EmbeddingModelChangedError as exc:
        report_log(
            reporter,
            f"[*] {exc} Re-embedding existing chunks with the new embedding model...",
            stage="kb-build",
            level="warning",
        )
        index, meta = rebuild_index_from_doc_store(
            doc_store_path,
            index_path,
            meta_path,
            embed_model,
            model_client,
            np,
            faiss,
            embedding_cache,
            total_records=len(existing_chunk_ids),
            reporter=reporter,
        )
    except CorruptIndexError as exc:
        report_log(
            reporter,
            f"[!] Existing FAISS index is unreadable; rebuilding it from doc_store.jsonl. {exc}",
            stage="kb-build",
            level="warning",
        )
        index, meta = rebuild_index_from_doc_store(
            doc_store_path,
            index_path,
            meta_path,
            embed_model,
            model_client,
            np,
            faiss,
            embedding_cache,
            total_records=len(existing_chunk_ids),
            reporter=reporter,
        )

    # Seed the durable vector cache from a readable legacy index once. This is
    # local work and avoids an API bill during the first migration.
    if (
        index is not None
        and existing_chunk_ids
        and embedding_cache.seeded_index_size() != int(index.ntotal)
    ):
        report_log(
            reporter,
            "[*] Seeding the embedding cache from the existing FAISS index...",
            stage="kb-build",
        )
        backfill_embedding_cache_from_index(
            doc_store_path,
            index,
            embedding_cache,
            np,
            reporter=reporter,
        )
        embedding_cache.mark_index_seeded(index.ntotal)

    if index is not None:
        stored_chunks = len(existing_chunk_ids)
        index_ids = get_index_vector_ids(index, faiss)
        index_mismatch = (
            index_ids != stored_vector_ids
            if index_ids is not None
            else int(index.ntotal) != stored_chunks
        )
        if index_mismatch:
            report_log(
                reporter,
                f"[!] FAISS/doc_store mismatch: index has {int(index.ntotal)} vectors, doc_store has {stored_chunks} chunks. Rebuilding from the local embedding cache.",
                stage="kb-build",
                level="warning",
            )
            index, meta = rebuild_index_from_doc_store(
                doc_store_path,
                index_path,
                meta_path,
                embed_model,
                model_client,
                np,
                faiss,
                embedding_cache,
                total_records=stored_chunks,
                reporter=reporter,
            )
        else:
            current_total = int(index.ntotal)
            if meta.get("total_vectors") != current_total:
                meta["total_vectors"] = current_total
                meta_dirty = True
    elif existing_chunk_ids:
        report_log(
            reporter,
            "[!] FAISS index is missing while doc_store.jsonl has records. Rebuilding index.",
            stage="kb-build",
            level="warning",
        )
        index, meta = rebuild_index_from_doc_store(
            doc_store_path,
            index_path,
            meta_path,
            embed_model,
            model_client,
            np,
            faiss,
            embedding_cache,
            total_records=len(existing_chunk_ids),
            reporter=reporter,
        )
    next_id = max(meta.get("next_id", 1), max_vector_id + 1)
    if meta.get("next_id") != next_id:
        meta["next_id"] = next_id
        meta_dirty = True

    pending_records = []
    total_new = 0
    processed_files = 0

    for file_path in iter_paper_files(source_dir):
        processed_files += 1
        emit(
            reporter,
            "kb_file",
            stage="kb-build",
            current=processed_files,
            file_name=os.path.basename(file_path),
            file_path=file_path,
        )

        file_signature = get_file_signature(file_path)
        manifest_key = os.path.abspath(file_path)
        manifest_entry = source_manifest.get("files", {}).get(manifest_key)
        if can_reuse_manifest_entry(manifest_entry, file_signature, existing_chunk_ids):
            emit(
                reporter,
                "kb_progress",
                stage="kb-build",
                processed_files=processed_files,
                skipped_files=1,
                added_chunks=total_new,
                total_vectors=meta.get("total_vectors"),
            )
            continue

        text, doc_meta = extract_text_and_meta(file_path)
        if not text.strip():
            continue

        sections = split_sections(text)
        doc_id = doc_meta["doc_id"]
        file_chunk_ids = []
        for section_index, (section_title, section_text) in enumerate(sections):
            chunks = split_into_chunks(section_text)
            for chunk_index, chunk in enumerate(chunks):
                chunk_id = f"{doc_id}:{section_index}:{chunk_index}"
                file_chunk_ids.append(chunk_id)
                if chunk_id in existing_chunk_ids:
                    continue

                record = {
                    "chunk_id": chunk_id,
                    "doc_id": doc_id,
                    "source_path": doc_meta["source_path"],
                    "file_name": doc_meta["file_name"],
                    "title": doc_meta.get("title", ""),
                    "authors": doc_meta.get("authors", ""),
                    "year": doc_meta.get("year", ""),
                    "date": doc_meta.get("date", ""),
                    "file_type": doc_meta["file_type"],
                    "section_title": section_title,
                    "chunk_index": chunk_index,
                    "chunk_total": len(chunks),
                    "text": chunk,
                }
                pending_records.append(record)

                if len(pending_records) >= 64:
                    added, index = flush_batch(
                        pending_records,
                        model_client,
                        index,
                        meta,
                        doc_store_path,
                        meta_path,
                        np,
                        embedding_cache,
                        chunk_store,
                    )
                    index_dirty = True
                    total_new += added
                    emit(reporter, "kb_progress", stage="kb-build", added_chunks=total_new, total_vectors=meta.get("total_vectors"))
                    for rec in pending_records:
                        existing_chunk_ids.add(rec["chunk_id"])
                    pending_records = []

        source_manifest.setdefault("files", {})[manifest_key] = {
            **file_signature,
            "doc_id": doc_id,
            "chunk_ids": file_chunk_ids,
            "parser_version": PARSER_VERSION,
            "chunker_version": CHUNKER_VERSION,
        }

    if pending_records:
        added, index = flush_batch(
            pending_records,
            model_client,
            index,
            meta,
            doc_store_path,
            meta_path,
            np,
            embedding_cache,
            chunk_store,
        )
        index_dirty = True
        total_new += added
        emit(reporter, "kb_progress", stage="kb-build", added_chunks=total_new, total_vectors=meta.get("total_vectors"))
        for rec in pending_records:
            existing_chunk_ids.add(rec["chunk_id"])

    if index is None:
        report_log(reporter, "[!] No vectors were created. Check source directory and parsers.", stage="kb-build")
        embedding_cache.close()
        chunk_store.close()
        build_lock.close()
        return

    if index_dirty:
        write_index_atomic(index, index_path, faiss)
        save_meta(meta_path, meta)
    elif meta_dirty:
        save_meta(meta_path, meta)

    if source_dir is not None:
        source_manifest["parser_version"] = PARSER_VERSION
        source_manifest["chunker_version"] = CHUNKER_VERSION
        save_source_manifest(manifest_path, source_manifest)
    if index is not None:
        embedding_cache.mark_index_seeded(index.ntotal)
    embedding_cache.close()
    chunk_store.close()
    build_lock.close()

    report_log(reporter, f"[+] KB build complete. Added {total_new} new chunks.", stage="kb-build", added_chunks=total_new)


def load_or_create_index(index_path, meta_path, embed_model):
    try:
        import faiss
    except Exception:
        raise

    if os.path.exists(index_path):
        if not os.path.exists(meta_path):
            raise CorruptIndexError("Index metadata is missing.")
        try:
            with open(meta_path, "r", encoding="utf-8") as f:
                meta = json.load(f)
        except Exception as exc:
            raise CorruptIndexError(f"Index metadata is unreadable: {exc}") from exc
        if meta.get("embedding_model") != embed_model:
            raise EmbeddingModelChangedError(meta.get("embedding_model"), embed_model)
        try:
            index = faiss.read_index(index_path)
        except Exception as exc:
            raise CorruptIndexError(f"{index_path}: {exc}") from exc
        return index, meta

    meta = {
        "embedding_model": embed_model,
        "dimension": None,
        "total_vectors": 0,
        "next_id": 1,
        "updated_at": None,
    }
    index = None
    return index, meta


def new_meta(embed_model):
    return {
        "embedding_model": embed_model,
        "dimension": None,
        "total_vectors": 0,
        "next_id": 1,
        "updated_at": None,
    }


def get_index_vector_ids(index, faiss):
    id_map = getattr(index, "id_map", None)
    if id_map is None:
        return None
    try:
        return {int(value) for value in faiss.vector_to_array(id_map)}
    except Exception:
        return None


def rebuild_index_from_doc_store(
    doc_store_path,
    index_path,
    meta_path,
    embed_model,
    model_client,
    np,
    faiss,
    embedding_cache,
    total_records=None,
    reporter=None,
):
    meta = new_meta(embed_model)
    if not os.path.exists(doc_store_path):
        report_log(
            reporter,
            "[!] doc_store.jsonl not found; starting a fresh KB index.",
            stage="kb-build",
            level="warning",
        )
        return None, meta

    index = None
    batch = []
    recovered = 0
    embedded = 0
    max_vector_id = 0
    total_label = str(total_records) if total_records is not None else "unknown"
    last_log = time.monotonic()
    report_log(
        reporter,
        f"[*] Rebuilding FAISS index from doc_store.jsonl: 0/{total_label} chunks recovered...",
        stage="kb-build",
    )

    for record in iter_doc_store_records(doc_store_path):
        batch.append(record)
        if len(batch) >= 64:
            index, added, batch_max_id, cache_misses = add_existing_records_to_index(
                batch,
                model_client,
                index,
                meta,
                np,
                faiss,
                embedding_cache,
            )
            recovered += added
            embedded += cache_misses
            max_vector_id = max(max_vector_id, batch_max_id)
            emit(reporter, "kb_rebuild_progress", stage="kb-build", recovered_chunks=recovered)
            if recovered % 2048 == 0 or time.monotonic() - last_log >= 30:
                report_log(
                    reporter,
                    f"[*] Rebuilding FAISS index from doc_store.jsonl: {recovered}/{total_label} chunks recovered...",
                    stage="kb-build",
                    recovered_chunks=recovered,
                    total_chunks=total_records,
                )
                last_log = time.monotonic()
            batch = []

    if batch:
        index, added, batch_max_id, cache_misses = add_existing_records_to_index(
            batch,
            model_client,
            index,
            meta,
            np,
            faiss,
            embedding_cache,
        )
        recovered += added
        embedded += cache_misses
        max_vector_id = max(max_vector_id, batch_max_id)
        emit(reporter, "kb_rebuild_progress", stage="kb-build", recovered_chunks=recovered)
        report_log(
            reporter,
            f"[*] Rebuilding FAISS index from doc_store.jsonl: {recovered}/{total_label} chunks recovered...",
            stage="kb-build",
            recovered_chunks=recovered,
            total_chunks=total_records,
        )

    if index is None:
        report_log(
            reporter,
            "[!] doc_store.jsonl has no usable records; starting a fresh KB index.",
            stage="kb-build",
            level="warning",
        )
        return None, meta

    meta["next_id"] = max_vector_id + 1
    meta["total_vectors"] = int(index.ntotal)
    meta["updated_at"] = datetime.datetime.now().isoformat()
    write_index_atomic(index, index_path, faiss)
    save_meta(meta_path, meta)
    report_log(
        reporter,
        f"[+] Rebuilt FAISS index from doc_store.jsonl. Recovered {recovered} chunks; embedded {embedded} cache misses.",
        stage="kb-build",
        recovered_chunks=recovered,
        embedded_chunks=embedded,
    )
    return index, meta


def iter_doc_store_records(doc_store_path):
    with open(doc_store_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                record = json.loads(line)
            except Exception:
                continue
            if not record.get("text"):
                continue
            vector_id = record.get("vector_id")
            if not isinstance(vector_id, int):
                continue
            yield record


def vectors_for_records(
    records,
    model_client,
    embedding_cache,
    np,
    expected_dimension=None,
):
    text_hashes = [chunk_text_hash(record["text"]) for record in records]
    vectors_by_hash = embedding_cache.get_many(text_hashes)
    if expected_dimension:
        vectors_by_hash = {
            key: vector
            for key, vector in vectors_by_hash.items()
            if vector.size == int(expected_dimension)
        }

    missing = {}
    for text_hash, record in zip(text_hashes, records):
        if text_hash not in vectors_by_hash:
            missing.setdefault(text_hash, record["text"])

    if missing:
        missing_keys = list(missing)
        vectors = model_client.embed_texts([missing[key] for key in missing_keys])
        vectors_np = normalize_vectors(np.asarray(vectors, dtype="float32"), np)
        if expected_dimension and vectors_np.shape[1] != int(expected_dimension):
            raise RuntimeError("Embedding dimension mismatch while filling the cache.")
        new_vectors = {
            key: vectors_np[offset]
            for offset, key in enumerate(missing_keys)
        }
        embedding_cache.put_many(new_vectors)
        vectors_by_hash.update(new_vectors)

    ordered = np.asarray(
        [vectors_by_hash[text_hash] for text_hash in text_hashes],
        dtype="float32",
    )
    return ordered, len(missing)


def backfill_embedding_cache_from_index(
    doc_store_path,
    index,
    embedding_cache,
    np,
    reporter=None,
):
    batch = []
    cached = 0
    last_log = time.monotonic()

    def flush(records):
        if not records:
            return 0
        text_hashes = [chunk_text_hash(record["text"]) for record in records]
        existing = embedding_cache.get_many(text_hashes)
        recovered = {}
        for record, text_hash in zip(records, text_hashes):
            if text_hash in existing or text_hash in recovered:
                continue
            try:
                vector = np.asarray(
                    index.reconstruct(int(record["vector_id"])),
                    dtype="float32",
                )
            except Exception:
                continue
            if vector.size != int(index.d):
                continue
            norm = float(np.linalg.norm(vector))
            if norm:
                vector = vector / norm
            recovered[text_hash] = vector
        embedding_cache.put_many(recovered)
        return len(recovered)

    for record in iter_doc_store_records(doc_store_path):
        batch.append(record)
        if len(batch) >= 512:
            cached += flush(batch)
            batch = []
            if cached and time.monotonic() - last_log >= 30:
                report_log(
                    reporter,
                    f"[*] Seeded {cached} embeddings from FAISS...",
                    stage="kb-build",
                    cached_embeddings=cached,
                )
                last_log = time.monotonic()
    cached += flush(batch)
    report_log(
        reporter,
        f"[+] Seeded {cached} embeddings from the existing FAISS index.",
        stage="kb-build",
        cached_embeddings=cached,
    )
    return cached


def add_existing_records_to_index(
    records, model_client, index, meta, np, faiss, embedding_cache
):
    vectors_np, cache_misses = vectors_for_records(
        records,
        model_client,
        embedding_cache,
        np,
        expected_dimension=meta.get("dimension"),
    )

    if index is None:
        dimension = vectors_np.shape[1]
        meta["dimension"] = dimension
        base_index = faiss.IndexFlatIP(dimension)
        index = faiss.IndexIDMap2(base_index)

    if meta.get("dimension") and vectors_np.shape[1] != meta["dimension"]:
        raise RuntimeError("Embedding dimension mismatch while rebuilding index.")

    vector_ids = [int(record["vector_id"]) for record in records]
    ids_np = np.array(vector_ids, dtype="int64")
    index.add_with_ids(vectors_np, ids_np)
    meta["total_vectors"] = int(index.ntotal)
    meta["updated_at"] = datetime.datetime.now().isoformat()
    return index, len(records), max(vector_ids), cache_misses


def flush_batch(
    records,
    model_client,
    index,
    meta,
    doc_store_path,
    meta_path,
    np,
    embedding_cache,
    chunk_store,
):
    vectors_np, _ = vectors_for_records(
        records,
        model_client,
        embedding_cache,
        np,
        expected_dimension=meta.get("dimension"),
    )

    if index is None:
        dimension = vectors_np.shape[1]
        meta["dimension"] = dimension
        try:
            import faiss
        except Exception:
            raise
        base_index = faiss.IndexFlatIP(dimension)
        index = faiss.IndexIDMap2(base_index)

    if meta.get("dimension") and vectors_np.shape[1] != meta["dimension"]:
        raise RuntimeError("Embedding dimension mismatch. Rebuild index to continue.")

    vector_ids = []
    for record in records:
        record["vector_id"] = meta["next_id"]
        meta["next_id"] += 1
        vector_ids.append(record["vector_id"])

    ids_np = np.array(vector_ids, dtype="int64")
    index.add_with_ids(vectors_np, ids_np)

    # The transactional SQLite store is canonical. JSONL remains as a
    # compatibility projection for the existing search code.
    chunk_store.put_records(records)
    with open(doc_store_path, "a", encoding="utf-8") as f:
        for record in records:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")
        f.flush()
        os.fsync(f.fileno())

    meta["total_vectors"] = int(index.ntotal)
    meta["updated_at"] = datetime.datetime.now().isoformat()
    save_meta(meta_path, meta)

    return len(records), index


def normalize_vectors(vectors, np):
    norms = np.linalg.norm(vectors, axis=1, keepdims=True)
    norms[norms == 0] = 1.0
    return vectors / norms


def load_existing_chunk_ids(doc_store_path):
    existing, max_vector_id, _ = inspect_doc_store(doc_store_path)
    return existing, max_vector_id


def inspect_doc_store(doc_store_path):
    existing = set()
    max_vector_id = 0
    invalid_lines = 0
    if not os.path.exists(doc_store_path):
        return existing, max_vector_id, invalid_lines
    with open(doc_store_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                item = json.loads(line)
            except Exception:
                invalid_lines += 1
                continue
            chunk_id = item.get("chunk_id")
            vector_id = item.get("vector_id")
            if not chunk_id or not item.get("text") or not isinstance(vector_id, int):
                invalid_lines += 1
                continue
            existing.add(chunk_id)
            max_vector_id = max(max_vector_id, vector_id)
    return existing, max_vector_id, invalid_lines


def load_source_manifest(manifest_path):
    if not os.path.exists(manifest_path):
        return {"files": {}}
    try:
        with open(manifest_path, "r", encoding="utf-8") as f:
            manifest = json.load(f)
    except Exception:
        return {"files": {}}
    if not isinstance(manifest, dict) or not isinstance(manifest.get("files"), dict):
        return {"files": {}}
    return manifest


def save_source_manifest(manifest_path, manifest):
    save_meta(manifest_path, manifest)


def get_file_signature(file_path):
    try:
        stat = os.stat(file_path)
    except Exception:
        return None
    return {
        "size": int(stat.st_size),
        "mtime_ns": int(stat.st_mtime_ns),
    }


def can_reuse_manifest_entry(entry, signature, existing_chunk_ids):
    if not entry or not signature:
        return False
    if entry.get("parser_version") != PARSER_VERSION:
        return False
    if entry.get("chunker_version") != CHUNKER_VERSION:
        return False
    if entry.get("size") != signature["size"]:
        return False
    if entry.get("mtime_ns") != signature["mtime_ns"]:
        return False
    chunk_ids = entry.get("chunk_ids")
    if not isinstance(chunk_ids, list):
        return False
    return all(chunk_id in existing_chunk_ids for chunk_id in chunk_ids)


def remove_sqlite_store(path):
    for candidate in (path, f"{path}-wal", f"{path}-shm"):
        if os.path.exists(candidate):
            os.remove(candidate)


def acquire_build_lock(kb_dir):
    lock_path = os.path.join(kb_dir, "build.lock")
    lock_file = open(lock_path, "a+b")
    if os.path.getsize(lock_path) == 0:
        lock_file.write(b"0")
        lock_file.flush()
    lock_file.seek(0)
    try:
        if os.name == "nt":
            import msvcrt

            msvcrt.locking(lock_file.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl

            fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError as exc:
        lock_file.close()
        raise RuntimeError("Another knowledge-base build is already running.") from exc
    return lock_file


def iter_paper_files(source_dir):
    if not source_dir:
        return
    if isinstance(source_dir, (list, tuple, set)):
        for item in source_dir:
            yield from iter_paper_files(item)
        return
    for root, _, files in os.walk(source_dir):
        for name in files:
            if name.endswith(SKIP_SUFFIXES):
                continue
            ext = os.path.splitext(name)[1].lower()
            if ext not in SUPPORTED_EXTS:
                continue
            yield os.path.join(root, name)


def extract_text_and_meta(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".doc":
        print(f"[!] Unsupported .doc file skipped: {os.path.basename(file_path)}")
        return "", {}
    try:
        stat = os.stat(file_path)
    except Exception as exc:
        remove_corrupt_file(file_path, f"stat failed: {exc}")
        return "", {}
    doc_id = hash_doc_id(file_path, stat.st_size, stat.st_mtime)
    date_tag = find_date_in_path(file_path)
    year = date_tag.split("-")[0] if date_tag else ""

    metadata = {
        "doc_id": doc_id,
        "source_path": file_path,
        "file_name": os.path.basename(file_path),
        "file_type": ext.lstrip("."),
        "date": date_tag or "",
        "year": year,
    }

    if ext == ".pdf":
        text, pdf_meta = extract_pdf_text(file_path)
        metadata["title"] = pdf_meta.get("title") or derive_title_from_filename(file_path)
        metadata["authors"] = pdf_meta.get("author") or ""
        return text, metadata

    if ext == ".docx":
        text = extract_docx_text(file_path)
        metadata["title"] = derive_title_from_text(text) or derive_title_from_filename(file_path)
        return text, metadata

    text = read_text_file(file_path)
    metadata["title"] = derive_title_from_text(text) or derive_title_from_filename(file_path)
    return text, metadata


def extract_pdf_text(file_path):
    doc = None
    try:
        with suppress_mupdf_errors():
            doc = fitz.open(file_path)
            meta = doc.metadata or {}
            pages = []
            for page in doc:
                try:
                    pages.append(extract_page_text(page))
                except Exception:
                    continue
        return "\n\n".join(pages), meta
    except Exception as exc:
        remove_corrupt_file(file_path, f"pdf parse failed: {exc}")
        return "", {}
    finally:
        if doc is not None:
            with contextlib.suppress(Exception):
                doc.close()


@contextlib.contextmanager
def suppress_mupdf_errors():
    tools = getattr(fitz, "TOOLS", None)
    if tools:
        for name in ("mupdf_display_errors", "mupdf_display_warnings"):
            func = getattr(tools, name, None)
            if callable(func):
                with contextlib.suppress(Exception):
                    func(False)
    with open(os.devnull, "w") as devnull:
        with contextlib.redirect_stderr(devnull):
            yield


def extract_page_text(page):
    blocks = page.get_text("blocks")
    if not blocks:
        return page.get_text()

    text_blocks = [b for b in blocks if b[4].strip()]
    if not text_blocks:
        return ""

    mid = page.rect.width / 2
    left = [b for b in text_blocks if (b[0] + b[2]) / 2 <= mid]
    right = [b for b in text_blocks if (b[0] + b[2]) / 2 > mid]

    if left and right:
        left_sorted = sorted(left, key=lambda b: (b[1], b[0]))
        right_sorted = sorted(right, key=lambda b: (b[1], b[0]))
        ordered = left_sorted + right_sorted
    else:
        ordered = sorted(text_blocks, key=lambda b: (b[1], b[0]))

    lines = [b[4].strip() for b in ordered if b[4].strip()]
    return "\n".join(lines)


def extract_docx_text(file_path):
    try:
        from docx import Document
    except Exception:
        print("[!] Missing dependency: python-docx. Skipping docx file.")
        return ""

    try:
        doc = Document(file_path)
    except Exception as exc:
        remove_corrupt_file(file_path, f"docx parse failed: {exc}")
        return ""
    parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        parts.append(text)
    return "\n\n".join(parts)


def read_text_file(file_path):
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception as exc:
            remove_corrupt_file(file_path, f"text read failed: {exc}")
            return ""
    except Exception as exc:
        remove_corrupt_file(file_path, f"text read failed: {exc}")
        return ""


def split_sections(text):
    lines = text.splitlines()
    sections = []
    current_title = "Body"
    current_lines = []

    for line in lines:
        stripped = line.strip()
        if is_heading_line(stripped):
            if current_lines:
                sections.append((current_title, "\n".join(current_lines).strip()))
            current_title = clean_heading(stripped)
            current_lines = []
        else:
            current_lines.append(line)

    if current_lines:
        sections.append((current_title, "\n".join(current_lines).strip()))

    return [(title, text) for title, text in sections if text]


def is_heading_line(line):
    if not line:
        return False
    if line.startswith("#"):
        return True
    lowered = line.lower()
    if lowered in {"abstract", "introduction", "methods", "method", "results", "conclusion", "references"}:
        return True
    if re.match(r"^\d+(\.\d+)*\s+\S+", line) and len(line) <= 80:
        return True
    if line.isupper() and 4 <= len(line) <= 80:
        return True
    if line.endswith(":") and len(line) <= 80:
        return True
    return False


def clean_heading(line):
    line = line.lstrip("#").strip()
    return line if line else "Section"


def split_into_chunks(text, min_chars=300, max_chars=600, overlap_ratio=0.15):
    if not text:
        return []

    overlap_chars = int(max_chars * overlap_ratio)
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    normalized_paragraphs = []
    for para in paragraphs:
        if len(para) <= max_chars:
            normalized_paragraphs.append(para)
        else:
            normalized_paragraphs.extend(split_long_text(para, max_chars, overlap_chars))

    chunks = []
    current = []
    current_len = 0

    for para in normalized_paragraphs:
        if current_len + len(para) + 2 > max_chars and current:
            chunk_text = "\n\n".join(current).strip()
            chunks.append(chunk_text)
            if overlap_chars > 0:
                overlap_text = chunk_text[-overlap_chars:]
                current = [overlap_text, para]
                current_len = len(overlap_text) + len(para)
            else:
                current = [para]
                current_len = len(para)
        else:
            current.append(para)
            current_len += len(para) + 2

    if current:
        chunks.append("\n\n".join(current).strip())

    if len(chunks) >= 2 and len(chunks[-1]) < min_chars:
        chunks[-2] = (chunks[-2] + "\n\n" + chunks[-1]).strip()
        chunks.pop()

    return [chunk for chunk in chunks if chunk]


def split_long_text(text, max_chars, overlap_chars):
    chunks = []
    start = 0
    step = max_chars - overlap_chars if max_chars > overlap_chars else max_chars
    while start < len(text):
        end = min(start + max_chars, len(text))
        chunks.append(text[start:end])
        if end >= len(text):
            break
        start += step
    return chunks


def derive_title_from_filename(file_path):
    base = os.path.splitext(os.path.basename(file_path))[0]
    base = re.sub(r"_\d{4}\.\d{4,5}(v\d+)?$", "", base)
    return base.replace("_", " ").strip()


def derive_title_from_text(text):
    if not text:
        return ""
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            return stripped.lstrip("#").strip()
        if stripped and len(stripped) <= 120:
            return stripped
    return ""


def hash_doc_id(path, size, mtime):
    payload = f"{path}|{size}|{int(mtime)}"
    return hashlib.sha1(payload.encode("utf-8")).hexdigest()


def find_date_in_path(path):
    parts = os.path.normpath(path).split(os.sep)
    for part in parts:
        if DATE_DIR_PATTERN.fullmatch(part):
            return part
    return ""


def save_meta(meta_path, meta):
    tmp_path = f"{meta_path}.{os.getpid()}.tmp"
    try:
        with open(tmp_path, "w", encoding="utf-8") as f:
            json.dump(meta, f, ensure_ascii=False, indent=2)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp_path, meta_path)
    finally:
        if os.path.exists(tmp_path):
            with contextlib.suppress(Exception):
                os.remove(tmp_path)


def write_index_atomic(index, index_path, faiss):
    tmp_path = f"{index_path}.{os.getpid()}.tmp"
    try:
        faiss.write_index(index, tmp_path)
        with open(tmp_path, "rb+") as f:
            os.fsync(f.fileno())
        os.replace(tmp_path, index_path)
    finally:
        if os.path.exists(tmp_path):
            with contextlib.suppress(Exception):
                os.remove(tmp_path)


def main():
    parser = argparse.ArgumentParser(description="Offline KB builder")
    parser.add_argument("--source", default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--kb-dir", default=DEFAULT_KB_DIR)
    parser.add_argument("--api", default=DEFAULT_API_KEY_PATH)
    parser.add_argument("--rebuild", action="store_true",
                        help="Wipe index, meta and doc_store.jsonl, then re-parse sources from scratch.")
    parser.add_argument("--refresh", action="store_true",
                        help="Keep parsed chunks (doc_store.jsonl) but re-embed them with the current "
                             "embedding model. Use after changing the embedding model.")
    args = parser.parse_args()

    build_kb(
        source_dir=args.source,
        kb_dir=args.kb_dir,
        api_key_path=args.api,
        rebuild=args.rebuild,
        refresh=args.refresh,
    )


if __name__ == "__main__":
    main()
